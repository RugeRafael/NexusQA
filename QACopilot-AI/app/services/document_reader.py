import re
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def read_document(file_path: str, file_content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext in ['.html', '.htm']:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(file_content, 'lxml')
            return soup.get_text(separator='\n', strip=True)
        elif ext == '.md':
            return file_content.decode('utf-8', errors='ignore')
        elif ext in ['.xlsx', '.xls']:
            import openpyxl
            import io
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"=== Hoja: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    row_text = ' | '.join([str(c) if c is not None else '' for c in row])
                    if row_text.strip('| '):
                        text_parts.append(row_text)
            return '\n'.join(text_parts)
        elif ext == '.docx':
            from docx import Document
            import io
            doc = Document(io.BytesIO(file_content))
            return '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == '.txt':
            return file_content.decode('utf-8', errors='ignore')
        else:
            return file_content.decode('utf-8', errors='ignore')
    except Exception as e:
        return f"Error leyendo documento: {str(e)}"


def parse_test_plan_html(file_content: bytes) -> dict:
    """
    Parsea el HTML del plan de pruebas.
    Soporta CP- y TC-, extrae módulo, sub-módulo y RF desde tc-meta-cell.
    """
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(file_content, 'lxml')

        result = {
            "proyecto": "",
            "qa": "",
            "fecha": "",
            "ambiente": "",
            "total_cps": 0,
            "rfs": [],
            "casos": []
        }

        # ── Metadata ──────────────────────────────────────────────
        header = soup.find(class_='header')
        if header:
            for item in header.find_all(class_='header-meta-item'):
                label_el = item.find(class_='label')
                value_el = item.find(class_='value')
                if label_el and value_el:
                    label = label_el.get_text(strip=True).lower()
                    value = value_el.get_text(strip=True)
                    if 'proyecto' in label: result['proyecto'] = value
                    elif 'qa' in label or 'engineer' in label: result['qa'] = value
                    elif 'fecha' in label: result['fecha'] = value
                    elif 'ambiente' in label: result['ambiente'] = value

        # ── RFs ───────────────────────────────────────────────────
        for section in soup.find_all(class_='section'):
            h2 = section.find('h2')
            if not h2: continue
            title = h2.get_text(strip=True)
            if 'Requerimientos' in title or 'Requisitos' in title:
                for row in section.find_all('tr'):
                    cells = row.find_all('td')
                    if len(cells) >= 2:
                        rf_id = cells[0].get_text(strip=True)
                        if re.match(r'RF-\d+', rf_id):
                            result['rfs'].append({
                                'id': rf_id,
                                'nombre': cells[1].get_text(strip=True),
                                'cantidad': cells[2].get_text(strip=True) if len(cells) > 2 else '',
                                'rango': cells[3].get_text(strip=True) if len(cells) > 3 else ''
                            })

        # ── Casos: estrategia 1 — tc-card (formato NexusQA/CRONOS) ─
        tc_cards = soup.find_all(class_='tc-card')
        if tc_cards:
            for card in tc_cards:
                head = card.find(class_='tc-head')
                if not head: continue

                tc_id_el = head.find(class_='tc-id') or head.find(class_='cp-id')
                tc_name_el = head.find(class_='tc-name') or head.find(class_='cp-name')
                priority_el = head.find(class_='priority-badge') or head.find(class_='priority-alta') or head.find(class_='priority-media') or head.find(class_='priority-baja')

                if not tc_id_el: continue
                cp_id = tc_id_el.get_text(strip=True)
                if not re.match(r'^(CP|TC)-\d+', cp_id): continue

                nombre = tc_name_el.get_text(strip=True) if tc_name_el else ''
                criticidad = priority_el.get_text(strip=True) if priority_el else 'Media'

                modulo = ''
                submodulo = ''
                rf = ''

                for cell in card.find_all(class_='tc-meta-cell'):
                    text = cell.get_text(strip=True)
                    # Remover prefijos de label
                    for prefix in ['Módulo', 'Modulo', 'Module']:
                        if text.startswith(prefix):
                            modulo = text[len(prefix):].strip()
                            break
                    for prefix in ['Sub-módulo', 'Sub-modulo', 'Submodulo', 'Submódulo']:
                        if text.startswith(prefix):
                            submodulo = text[len(prefix):].strip()
                            break
                    if text.startswith('Requerimiento'):
                        rf_candidate = text.replace('Requerimiento', '').strip()
                        if re.match(r'RF-\d+', rf_candidate):
                            rf = rf_candidate

                result['casos'].append({
                    'id': cp_id,
                    'nombre': nombre,
                    'rf': rf or 'RF-General',
                    'rf_nombre': '',
                    'modulo': modulo,
                    'submodulo': submodulo,
                    'criticidad': criticidad,
                    'marcado': False
                })

        # ── Casos: estrategia 2 — fallback regex en texto completo ─
        if not result['casos']:
            logger.info("Usando fallback regex en texto completo")
            seen_ids = set()
            for m in re.finditer(r'\b(CP|TC)-(\d{1,4})\b', soup.get_text()):
                cp_id = m.group(0)
                if cp_id not in seen_ids:
                    seen_ids.add(cp_id)
                    result['casos'].append({
                        'id': cp_id, 'nombre': cp_id,
                        'rf': 'RF-General', 'rf_nombre': '',
                        'modulo': '', 'submodulo': '',
                        'criticidad': 'Media', 'marcado': False
                    })

        result['total_cps'] = len(result['casos'])
        logger.info(f"Total casos: {result['total_cps']}, RFs: {len(result['rfs'])}")
        return result

    except Exception as e:
        logger.error(f"Error parsing test plan: {e}")
        return {"proyecto": "", "qa": "", "fecha": "", "ambiente": "",
                "total_cps": 0, "rfs": [], "casos": [], "error": str(e)}
